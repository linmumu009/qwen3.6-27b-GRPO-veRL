from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'reports/training_materials_20260916'
MD = ROOT / 'docs/training_materials_20260916'

def p(s): return ('p',s)
def h(s): return ('h',s)
def code(s): return ('code',s.strip())
def table(head, rows): return ('table',head,rows)

materials = [
('01_CPT语料与运行说明', 'CPT 语料与运行说明', [
[
p('本材料说明物流原文 CPT 的来源、输入格式和启动方法。运行示例对应 2026 年 9 月 11 日的完整知识语料：从 Qwen3.6-27B 的开源续训 Step120 出发，训练一遍。'),
h('一 原始语料来源'),
p('原始材料由用户提供，共 13 份：12 份 PDF、4,050 页，以及 1 个 HTML 教材包。正文经 MinerU 或本地文本层提取，保留章节、页码及来源关系。'),
table(['内容方向','原始材料'],[
['港口与运输','Ports and Waterways；Container Terminals and Automated Transport Systems'],
['配送与仓储','Distribution Logistics；Warehouse & Distribution Science'],
['供应链管理','JSI Supply Chain Manager’s Handbook；Managing Supply Chain Risk；MITx SCM Key Concepts'],
['交通与统计','Traffic Flow Theory；Transport Statistics Glossary'],
['贸易与战略','International Trade: Theory and Policy；Fundamentals of Global Strategy；Challenges and Opportunities in International Business'],
['仿真','Beyond Lean: Simulation in Practice'],
]),
p('在教材正文基础上，追加 205 条经核对的知识单元，来源包括 UNECE、FAA、UNCTAD、IFRS、IATA 及百科等公开资料。补充内容保留适用范围与出处。'),
],
[
h('二 CPT 语料形态'),
p('训练输入是连续知识文本，不是问答对。清洗后按完整文本块组织记录，去除无关版面内容并去重；表格保留结构，公式保留表达。正文不套聊天模板，不跨书混拼。'),
table(['语料版本','记录数','正文 token'],[
['全部 13 份教材','4,707','1,844,881'],
['新增知识单元','205','23,131'],
['合并训练集','4,912','1,868,012'],
]),
p('交付为同内容 JSONL 和 Parquet。训练器读取 text 字段；来源、章节及页码通过配套索引追踪。以下为字段示意，正文为演示文本。'),
code('''{"text":"仓储知识示例：库存盘点需要核对账面数量与实物数量。"}'''),
table(['读取规则','设置'],[
['数据入口','train.parquet；text_key=text'],
['结束标记','加载器追加一次 EOS，不在正文手工添加'],
['最大长度','4,096 token；本批最大 4,085'],
['超长处理','报错，不静默截断'],
['训练方式','原文下一 token 预测；无聊天模板'],
]),
p('本版本包含评测来源教材，属于定向学习实验。评测题目和答案文件本身不进入该语料；相关基准成绩不能作为独立泛化证明。'),
],
[
h('三 CPT 脚本运行指令'),
p('在 5 号机 llin-verl-trainer-m05-20260730 容器内执行。需已有 Step120 模型、分布式权重、固定 Megatron Bridge 及训练文件。脚本固定本批 4,912 条、batch 8、614 步，不能直接用于其他规模语料。'),
code('''source /usr/local/Ascend/ascend-toolkit/set_env.sh
cd /workspace/llin-verl-grpo
export PROJECT_ROOT="$PWD"
export TRAIN_FILE="$PWD/runs/llin-knowledge-complete-20260911/train.parquet"
export EXPECTED_CONTENT_TOKENS=1868012
sha=6ffa16684e617f657293918bf1d89ae8930
sha+=ec24e0e08da346f175c4710c1f1c8
export EXPECTED_TRAIN_SHA="$sha"
export RUN_NAME="llin-cpt-knowledge-replay-$(date +%Y%m%d-%H%M%S)"
export OUTPUT_DIR="$PWD/runs/$RUN_NAME"
bash scripts/run_cpt_knowledge_complete_one_epoch.sh'''),
p('默认起点为 llin-step120-opensource-20260825-02。16 张 NPU，TP4/PP2/CP2；学习率 5e-7 至 1e-7。第 614 步保存检查点，导出后用于推理。'),
p('此处复用 veRL 的 sft_trainer 入口，通过原文加载器执行 CPT。新输出目录须不存在；脚本自动检查文件指纹与完整加载。'),
h('依据'),
p('本仓库 docs/cpt_complete_package_20260911.md、docs/cpt_knowledge_complete_one_epoch_20260911.md；scripts/run_cpt_knowledge_complete_one_epoch.sh、scripts/qwen36_causal_lm_dataset.py。'),
]
]),
('02_通用开源SFT数据与运行说明','通用开源 SFT 数据与运行说明',[
[
p('本材料说明开源指令问答数据的提取、格式转换和 SFT 启动方法。数据为 20260702 版 25,167 条问答；训练入口采用现存的 Qwen3.6-27B MindSpeed-MM 全参脚本。'),
h('一 QA 提取流程'),
p('从已下载开源数据读取题目、参考解答和推理过程，先统一为四字段，再转换为对话。这里是已有标注的字段提取，不调用模型重新生成答案。'),
table(['数据来源','历史合并条数'],[
['MATH','12,500'],['GSM8K','7,473'],['Omni-MATH','4,428'],['AMO-Bench','50'],['PHYBench','200'],['ATLAS','256'],['C-Eval-dev','260'],['合计','25,167']]),
p('MATH 从 problem 取题目，从 solution 提取推理与 boxed 答案；GSM8K 从 question 取题目，按 answer 中的 #### 分离推导和最终答案。其他数据由各自转换器处理。'),
p('convert.py 根据配置中的 sft_selected 选择数据，按推理字段是否为空分流。历史合并集与当前转换配置分别管理，不能假设重新运行后仍得到完全相同的条数。'),
],
[
h('二 SFT 的训练数据形态'),
p('中间层每行一条 JSON，四字段分别是系统要求、题目、推理和参考答案。以下使用演示题说明结构。'),
code('''{
  "system_prompt": "请解答数学问题。",
  "task_input": "求解 2x + 3 = 7。",
  "thinking": "两边减去 3，再除以 2。",
  "golden_answer": "x = 2"
}'''),
p('训练层为 OpenAI messages JSONL。有推理时，reasoning_content 保存推理，content 保存带 answer 标签的答案；无推理时，content 直接保存答案。'),
code('''{"messages": [
  {"role":"system","content":"请解答数学问题。"},
  {"role":"user","content":"求解 2x + 3 = 7。"},
  {"role":"assistant",
   "reasoning_content":"两边减去 3，再除以 2。",
   "content":"<answer>\\nx = 2\\n</answer>"}
]}'''),
p('MindSpeed-MM 的 Qwen3.6 模板负责合并推理。历史 ms-swift 另有 reasoning_merged 版本，将推理写入 content 的 think 块；该文件共 25,167 条，其中 25,153 条包含 think 块。两种文件不能不检查模板就互换。'),
p('模型学习 assistant 的推理与答案，system 和 user 作为上下文。本路线没有沙箱工具轨迹，也不等同于 Step100 后的开源数据 GRPO 续训。'),
],
[
h('三 SFT 脚本运行指令'),
p('数据转换在 5 号机 huawei_train 项目执行。当前目录已改为按能力分类；下面以 Reasoning 目录作为新一批转换输入。该命令不重建历史七数据集配比。'),
code('''cd /data/renjunxiang/coding/huawei_train
python3 scripts/data/convert.py --list
python3 scripts/data/to_openai.py \\
  --input datasets/open_source/processed/with_golden_thinking/Reasoning \\
  --output datasets/open_source/sft/final/reasoning_new.jsonl'''),
p('复用历史 25,167 条时，在宿主机将冻结的原始 OpenAI 文件复制到训练容器可见的 /data 路径；不要用 reasoning_merged 替换它。'),
code('''src=/data3/llin/sft/datasets/open_source/huawei_train
dst=/data/renjunxiang/coding/huawei_train/datasets/open_source/sft/final
mkdir -p "$dst"
cp -n "$src/20260702_openai.jsonl" "$dst/20260702_openai.jsonl"
sha256sum "$src/20260702_openai.jsonl" "$dst/20260702_openai.jsonl"'''),
p('两文件指纹必须一致。启动前，在 scripts/train/qwen3_6_27B_sft.yaml 中确认 dataset 指向上述冻结文件，并把 training.save 改为新的输出目录。模型需要 HF 路径 /data/models/Qwen3.6-27B 和 DCP 路径 /data/models/Qwen3.6-27B-dcp。'),
],
[
h('三 SFT 脚本运行指令 续'),
p('在 5 号机宿主机启动现有脚本。脚本自行设置 Ascend 环境并进入 MindSpeed-MM 工程。此处整理启动方法，不将该现存配置描述为已完成的正式全量训练。'),
code('''docker exec -it mindspeed_mm_rjx bash \\
  /data/renjunxiang/coding/huawei_train/scripts/train/finetune_qwen3_6_27B_sft.sh'''),
table(['现存全参配置','值'],[
['模型与框架','Qwen3.6-27B；MindSpeed-MM FSDP2'],
['设备','单机 16 张 Ascend NPU'],
['最大序列长度','2,048'],
['micro batch / 梯度累积','1 / 1'],
['训练步数 / 保存间隔','500 / 100'],
['学习率','1e-5'],
]),
p('另有已保存的 ms-swift LoRA 实验：8 卡 FSDP2、rank 8、alpha 32、学习率 1e-4；后段从第 300 步 adapter 初始化，长度 4,096，追加 700 步。它不属于本页全参启动脚本，也不是加载完整训练状态的断点续训。'),
h('依据'),
p('5 号机 huawei_train：scripts/data/convert.py、converters/math.py、converters/gsm8k.py、to_openai.py；scripts/train/finetune_qwen3_6_27B_sft.sh 及配套 YAML。以上入口与配置已于 2026 年 9 月 16 日只读核对。'),
p('七数据集历史配比见本仓库 技术报告/sources/engineering_report_20260913_工程复现报告.md；reasoning_merged 条数已与服务器冻结文件核对。'),
]
]),
('03_沙箱AgenticSFT数据与运行说明','沙箱 Agentic SFT 数据与运行说明',[
[
p('本材料对应已完成的 Qwen3.5-9B 沙箱轨迹 SFT：2,028 条 16K 内训练候选，150 步。它与 Qwen3.6-27B 的 16 条纠错诊断 SFT 是不同实验。'),
h('一 数据构成'),
p('数据来自 PI agent 在物流沙箱中的多轮交互。环境包含物流数据库、业务文档及任务；轨迹保留问题、模型思考、工具调用、执行结果和最终回答。'),
table(['数据层','数量与范围'],[
['原始 OpenAI 轨迹副本','24 个 JSONL，共 36,000 条'],
['训练候选合并文件','2,028 条，完整轨迹不超过 16,384 token'],
['其中 SQL 结果强验证','27 条'],
['其中弱证据待复核','2,001 条'],
['隔离数据','250 条 validation/test；1,981 条 16–32K 长轨迹未混入上述文件'],
]),
p('处理链路：读取完整事件消息 → 对齐任务身份 → 检查工具调用与返回闭环 → 按判分和证据筛选 → 按任务切分 → 按真实 token 长度筛选 → 合并训练文件。'),
p('2,028 条是历史实际使用的候选集，其中 27 条达到当时的强验证标准；不能把整批数据表述为已逐条完成高质量审核。训练完成与数据质量是两个不同结论。'),
],
[
h('二 数据形态'),
p('原始层使用 OpenAI messages，包含 assistant.reasoning_content、tool_calls 和 tool 消息。适配层由 prepare_trajectory_sft.py 转为 ms-swift 的 assistant、tool_call、tool_response 角色。'),
table(['字段或消息','作用'],[
['messages','按真实顺序保存完整交互'],
['tools','bash、read、write、edit 的工具定义；实际保存为 JSON 字符串'],
['assistant','思考与回答；参与训练'],
['tool_call','工具名和参数；参与训练'],
['tool_response','工具执行结果；只提供上下文'],
]),
p('以下仅展示角色与字段，内容为演示；真实 tools 需包含完整参数定义。'),
code('''{"tools":"[完整工具 schema]", "messages":[
  {"role":"user","content":"读取业务说明并回答问题。"},
  {"role":"assistant","content":"先查阅说明。", "loss":true},
  {"role":"tool_call",
   "content":"{\\"name\\":\\"read\\",\\"arguments\\":{\\"path\\":\\"rules.md\\"}}",
   "loss":true},
  {"role":"tool_response","content":"业务说明正文", "loss":false},
  {"role":"assistant","content":"依据说明得出的答案。", "loss":true}
]}'''),
p('训练保留整条工具交互。推理写入 think 块，工具返回不作为模型应生成的答案。超出 16K 的记录不混入本次训练。'),
],
[
h('三 训练脚本运行指令'),
p('在 5 号机 qwen3.5-9b-msswift-trajory-sft 容器内执行。宿主机 /data3/llin/trajory_sft 对应容器 /workspace/sft。需使用该项目已验证的 ms-swift、Megatron 和 MindSpeed 环境。'),
code('''source /usr/local/Ascend/ascend-toolkit/set_env.sh
cd /workspace/sft
export MODEL=/models/Qwen3.5-9B
data_dir="$PWD/data/screened/upstream_correct_screen_v1"
export DATASET="$data_dir/train_candidates_16k_2028.jsonl"
export RUN_DIR="$PWD/runs/agentic-sft-$(date +%Y%m%d-%H%M%S)"
export TRAIN_ITERS=150
export SAVE_STEPS=15
export SAVE_TOTAL_LIMIT=10
export NO_SAVE_OPTIM=false
export NO_SAVE_RNG=false
bash scripts/run_qwen35_megatron_tp4_pp2_dp2.sh'''),
p('脚本默认只跑两步测试，因此正式复现必须显式设置 TRAIN_ITERS=150。既定参数为 16 张 NPU、TP4/PP2/DP2、micro batch 1、global batch 16、长度 16,384、学习率 1e-6 至 1e-7。'),
p('历史运行已完成 150 步并导出 10 个 HF 模型。本启动脚本保存 MCore 分布式检查点，推理前需使用对应导出流程；视觉模块不参与本次文本轨迹训练。'),
h('依据'),
p('相邻项目 qwen9b-trajory-SFT：scripts/prepare_trajectory_sft.py、scripts/run_qwen35_megatron_tp4_pp2_dp2.sh；updates 中 v0.5.1 数据合并、v0.6.0 训练配置及 v0.7.0 完成记录。'),
]
]),
('04_GRPO沙箱数据与运行说明','GRPO 沙箱数据与运行说明',[
[
p('本材料说明 Qwen3.6-27B 的双机沙箱 GRPO 路线。以 v15 DWH 修正版数据和当前 100 步入口说明运行方法；历史 Step100 训练使用 237 条，修正版为 236 条，不混写。'),
h('一 RL 沙箱数据构成'),
p('RL 数据由任务、沙箱环境和评分依据组成。模型在沙箱中实时生成多轮轨迹，同一道任务采样 4 条，再按组内相对奖励更新。历史源轨迹保存在参考文件中，不作为 GRPO 的答案输入。'),
table(['组成','内容'],[
['任务','v15 DWH 物流数据查询与分析指令'],
['环境','同源 logistics.sqlite、表结构及工作区'],
['评分依据','参考 SQL、结构化预期结果、答案类型和所需表字段'],
['在线交互','bash、read、edit、write；每条轨迹独立工作区'],
]),
table(['数据版本','train','val','test'],[
['历史 full277','237','20','20'],['修正 full276','236','20','20']]),
p('修正版删除了一条相同指令绑定冲突答案的训练任务。276 条参考 SQL 均已通过同源数据库执行及结果匹配，但这不等于全部业务题意已完成人工确认。'),
],
[
h('二 数据形态'),
p('训练器读取 Parquet。prompt 只包含系统要求和用户任务；评分所需答案与 SQL 放在 reward_model.ground_truth 中，不拼入模型可见问题。'),
table(['字段','含义'],[
['data_source / agent_name','数据来源标记；pi_agent 执行循环'],
['prompt','system 与 user 消息'],
['reward_model.\nground_truth','任务、环境、答案类型、预期结果及参考 SQL'],
['extra_info','split、来源指纹、工具选择和环境创建参数'],
]),
p('以下为关键字段示意，省略完整工具参数和来源字段。'),
code('''{
  "data_source":"boss_pi_aligned_v1",
  "agent_name":"pi_agent",
  "prompt":[
    {"role":"system","content":"沙箱操作要求"},
    {"role":"user","content":"物流查询任务"}
  ],
  "reward_model":{"style":"rule","ground_truth":{
    "task_id":"任务编号", "environment_id":"环境编号",
    "answer_type":"答案类型",
    "expected_value_json":"序列化预期结果",
    "verification_sql":"经核验的只读 SQL"
  }},
  "extra_info":{"split":"train"}
}'''),
p('SFT 学习已有回答；本路线则在 rollout 阶段生成新的回答与工具操作，再由评分器给出奖励。两个阶段的数据文件不能直接互换。'),
],
[
h('三 脚本运行指令'),
p('5 号机负责训练，16 卡 TP4/PP2/CP2；6 号机负责 rollout，16 卡 TP8/DP2。两侧需同步项目、模型和沙箱，使用既定 veRL 环境；在专用空闲环境启动 Ray。'),
code('''# 5 号机训练容器
cd /workspace/llin-verl-grpo
bash scripts/start_ray_m05.sh

# 6 号机 rollout 容器
cd /workspace/llin-verl-grpo
bash scripts/start_ray_m06.sh'''),
p('Ray 两个角色就绪后，在 5 号机训练容器执行修正版数据检查与启动。'),
code('''cd /workspace/llin-verl-grpo
export PROJECT_ROOT="$PWD"
export PYTHONPATH="$PWD/runtime:$PWD:${PYTHONPATH:-}"
export DATA_DIR="$PWD/data/boss_v15_dwh_full276_20260806/dataset"
python3 scripts/check_boss_alignment_contract.py \\
  --data-dir "$DATA_DIR"
export RUN_NAME="llin-grpo-full276-$(date +%Y%m%d-%H%M%S)"
bash scripts/launch_pi_formal_100step_12groups.sh'''),
p('该入口从原始 /models/Qwen3.6-27B 训练 100 步。每步更新 4 组任务，每组 4 条轨迹；12 组指可并行在途组数，不是每步更新量。脚本同时检查远端数据并验证最终检查点。'),
p('本命令使用修正版 full276，不复刻旧 full277。8 月 25 日的开源续训 Step120 属于另一条 GRPO 路线，不在此启动。'),
h('依据'),
p('本仓库 scripts/prepare_boss_aligned_dataset.py、launch_pi_formal_100step_12groups.sh、run_pi_formal_100step_12groups.sh；docs/training_data_provenance_quality_audit_20260806.artifact.json。'),
]
])
]

def fonts(style, size, east='宋体'):
    style.font.name='Times New Roman'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr().get_or_add_rFonts()
    for k in list(rf.attrib):
        if 'Theme' in k:del rf.attrib[k]
    for k in ('ascii','hAnsi','cs'): rf.set(qn('w:'+k),'Times New Roman')
    rf.set(qn('w:eastAsia'),east)

for name,title,pages in materials:
    d=Document(); sec=d.sections[0]
    for st in d.styles:
        for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for s,sz,east in [('Normal',15,'宋体'),('Title',24,'黑体'),('Heading 1',18,'黑体'),('Heading 2',16,'黑体')]:
        st=d.styles[s];fonts(st,sz,east);f=st.paragraph_format;f.line_spacing=1.5;f.space_after=Pt(3);f.space_before=Pt(0)
        f.first_line_indent=Pt(30 if s=='Normal' else 0)
        ind=st.element.get_or_add_pPr().find(qn('w:ind'))
        if ind is not None:ind.set(qn('w:firstLineChars'),'200' if s=='Normal' else '0')
    st=d.styles.add_style('CodeText',1);fonts(st,12);st.paragraph_format.line_spacing=1.5;st.paragraph_format.space_after=Pt(0);st.paragraph_format.first_line_indent=Pt(0)
    d.add_paragraph(title,'Title')
    q=d.add_paragraph('数据与训练方法说明  2026年9月16日');q.paragraph_format.first_line_indent=Pt(0)
    q._p.get_or_add_pPr().find(qn('w:ind')).set(qn('w:firstLineChars'),'0')
    md=['# '+title,'','数据与训练方法说明  2026年9月16日','']
    for ix,page in enumerate(pages):
        # Let sections flow naturally so short carry-over paragraphs do not create sparse pages.
        source_mode=False
        for b in page:
            if b[0] in ('p','h'):
                if b[0]=='h':source_mode=b[1]=='依据'
                q=d.add_paragraph(b[1], 'Heading 1' if b[0]=='h' else 'Normal')
                if source_mode and b[0]=='p':
                    for r in q.runs:r.font.size=Pt(12)
                    q.paragraph_format.line_spacing=1.2
                md += [('## ' if b[0]=='h' else '')+b[1],'']
            elif b[0]=='code':
                for li,line in enumerate(b[1].splitlines()):
                    q=d.add_paragraph(line,'CodeText')
                    q.paragraph_format.keep_with_next=li<len(b[1].splitlines())-1
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');q._p.get_or_add_pPr().append(sh)
                md += ['```',b[1],'```','']
            else:
                t=d.add_table(rows=1,cols=len(b[1]));t.autofit=False
                widths=[6,11] if len(b[1])==2 else ([9,4,4] if len(b[1])==3 else [8,3,3,3])
                for j,col in enumerate(t.columns):col.width=Cm(widths[j])
                for j,x in enumerate(b[1]):t.rows[0].cells[j].text=x
                for row in b[2]:
                    c=t.add_row().cells
                    for j,x in enumerate(row):c[j].text=str(x)
                for ri,row in enumerate(t.rows):
                    for j,c in enumerate(row.cells):
                        c.width=Cm(widths[j]);c.vertical_alignment=1
                        for q in c.paragraphs:
                            q.paragraph_format.first_line_indent=Pt(0);q.paragraph_format.space_after=Pt(1);q.paragraph_format.space_before=Pt(1)
                            q._p.get_or_add_pPr().find(qn('w:ind')).set(qn('w:firstLineChars'),'0')
                            for r in q.runs:r.bold=(ri==0)
                            if ri==0:q.paragraph_format.keep_with_next=True
                        if ri==0:
                            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E7E7');c._tc.get_or_add_tcPr().append(sh)
                    trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                md += ['| '+' | '.join(b[1])+' |','| '+' | '.join(['---']*len(b[1]))+' |']+['| '+' | '.join(x.replace('\n','<br>') for x in r)+' |' for r in b[2]]+['']
    foot=sec.footer.paragraphs[0];foot.alignment=2;foot.paragraph_format.first_line_indent=Pt(0)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    d.core_properties.author='';d.core_properties.title=title
    d.save(OUT/(name+'.docx'));(MD/(name+'.md')).write_text('\n'.join(md),encoding='utf-8')
print('Created',len(materials),'documents')
